"""Accountant packet: Markdown + CSV (+ XLSX when openpyxl is available)."""
from __future__ import annotations

import csv
from pathlib import Path

from ..util import money, now_iso, plain
from . import ledger as ledger_service

PACKET_ID = "pkt_2025_demo"

OPEN_QUESTIONS = [
    "Should the canning line be expensed under §179, bonus depreciation, or depreciated?",
    "How should the solar and EV-charger assets be treated for NY vs federal purposes?",
    "Are taproom tips correctly separated from service charges in the payroll records?",
    "Does the IT-212 investment credit interact with any federal credit already claimed?",
    "What documentation does the accountant require for the two veteran hires?",
]


def _markdown(store, summary: dict, credits: list[dict]) -> str:
    rows = ledger_service.transactions(store.ledger)
    lines = [
        "# Accountant Packet — Gowanus Kettle Brewing Co. LLC",
        "",
        "> **Synthetic demo — not tax advice. Prepared for professional review, not filed.**",
        "",
        f"Tax year: **{summary['year']}** · Generated: {now_iso()} · Transactions: {summary['transaction_count']}",
        "",
        "## Summary",
        "",
        "| Line | Amount |",
        "|---|---:|",
        f"| Revenue | {money(summary['revenue_cents'])} |",
        f"| Expenses | {money(summary['expense_cents'])} |",
        f"| Capital assets | {money(summary['capital_cents'])} |",
        f"| Payroll | {money(summary['payroll_cents'])} |",
        f"| Tips | {money(summary['tips_cents'])} |",
        f"| Net | {money(summary['net_cents'])} |",
        f"| Reconciled delta | {money(summary['reconciled_delta_cents'])} |",
        f"| Missing receipts | {summary['missing_receipts']} |",
        f"| Duplicate candidates merged | {summary['duplicates']} |",
        "",
        "## Ledger",
        "",
        "| Date | Vendor | Description | Amount | Category | Source email | Flag |",
        "|---|---|---|---:|---|---|---|",
    ]
    for r in rows:
        flag = "missing receipt" if r.get("needs_receipt") else ("duplicate" if r.get("duplicate_of") else "")
        lines.append(
            f"| {r['date']} | {r['vendor']} | {r['description'][:48]} | {money(r['amount_cents'])} | "
            f"{r['category']} | {r.get('source_message_id', '—')} | {flag} |"
        )
    lines += ["", "## Credits in process", "", "| Form | Credit | Amount | Status |", "|---|---|---:|---|"]
    for c in credits:
        lines.append(f"| {c['form']} | {c['program']} | {money(c['amount_estimate_cents'])} | {c['status']} |")
    lines += ["", "## Evidence index", ""]
    for e in ledger_service.evidence_index(store.ledger)[:40]:
        lines.append(f"- `{e['transaction_id']}` → {e['source_message_id']} — “{e['source_quote'][:120]}”")
    lines += ["", "## Open questions for the accountant", ""]
    lines += [f"{i}. {q}" for i, q in enumerate(OPEN_QUESTIONS, 1)]
    lines += [
        "",
        "## Limitations",
        "",
        "- Synthetic demo data; not a real business, filing or tax position.",
        "- Email evidence only; no bank feed or audited statements.",
        "- Not tax advice. The accountant remains the professional.",
        "",
    ]
    return "\n".join(lines)


def _csv(path: Path, rows: list[dict]) -> None:
    with path.open("w", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["date", "vendor", "description", "amount", "direction", "category",
                         "source_message_id", "confidence", "needs_receipt", "duplicate_of"])
        for r in rows:
            writer.writerow([r["date"], r["vendor"], r["description"], plain(r["amount_cents"]),
                             r.get("direction", "purchase"), r["category"], r.get("source_message_id", ""),
                             r.get("confidence", ""), r.get("needs_receipt", False), r.get("duplicate_of") or ""])


def _xlsx(path: Path, summary: dict, rows: list[dict], credits: list[dict]) -> bool:
    try:
        from openpyxl import Workbook
    except Exception:  # noqa: BLE001 - optional dependency
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Line", "Amount (USD)"])
    for label, key in [("Revenue", "revenue_cents"), ("Expenses", "expense_cents"),
                       ("Capital assets", "capital_cents"), ("Payroll", "payroll_cents"),
                       ("Tips", "tips_cents"), ("Net", "net_cents")]:
        ws.append([label, round(summary[key] / 100, 2)])
    ws.append(["Reconciled delta", round(summary["reconciled_delta_cents"] / 100, 2)])
    ws.append(["Missing receipts", summary["missing_receipts"]])
    ledger_ws = wb.create_sheet("Ledger")
    ledger_ws.append(["Date", "Vendor", "Description", "Amount", "Category", "Source", "Flag"])
    for r in rows:
        ledger_ws.append([r["date"], r["vendor"], r["description"], round(r["amount_cents"] / 100, 2),
                          r["category"], r.get("source_message_id"), "missing receipt" if r.get("needs_receipt") else ""])
    credits_ws = wb.create_sheet("Credits")
    credits_ws.append(["Form", "Program", "Amount", "Status"])
    for c in credits:
        credits_ws.append([c["form"], c["program"], round(c["amount_estimate_cents"] / 100, 2), c["status"]])
    wb.save(path)
    return True


def build(store) -> dict:
    summary = ledger_service.summary(store.ledger)
    rows = ledger_service.transactions(store.ledger)
    credits = [{"form": a.get("form"), "program": a.get("program"),
                "amount_estimate_cents": a.get("amount_estimate_cents", 0), "status": a.get("status")}
               for a in store.applications]
    out = Path(store.data) / "packets" / PACKET_ID
    out.mkdir(parents=True, exist_ok=True)
    md = _markdown(store, summary, credits)
    (out / "review-package.md").write_text(md)
    _csv(out / "ledger.csv", rows)
    files = [
        {"name": "review-package.md", "kind": "markdown", "path": str(out / "review-package.md")},
        {"name": "ledger.csv", "kind": "csv", "path": str(out / "ledger.csv")},
    ]
    if _xlsx(out / "summary.xlsx", summary, rows, credits):
        files.append({"name": "summary.xlsx", "kind": "xlsx", "path": str(out / "summary.xlsx")})
    packet = {
        "id": PACKET_ID,
        "year": summary["year"],
        "generated_at": now_iso(),
        "summary": summary,
        "files": files,
        "open_questions": OPEN_QUESTIONS,
        "markdown": md,
        "status": "prepared_for_review",
        "disclaimer": "Synthetic demo — not tax advice. Prepared for professional review, not filed.",
    }
    store.packet = packet
    return packet
